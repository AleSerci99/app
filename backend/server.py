from fastapi import FastAPI, APIRouter, HTTPException, Depends, status, UploadFile, File, Query
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from fastapi.responses import Response
from fastapi.concurrency import run_in_threadpool
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
import os
import logging
import base64
import io
import requests
from pathlib import Path
from pydantic import BaseModel, EmailStr, Field
from typing import List, Optional, Dict
import uuid
from datetime import datetime, timezone, date

import jwt
import bcrypt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet


ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

# MongoDB connection
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

JWT_SECRET = os.environ['JWT_SECRET']
JWT_ALGO = "HS256"

# --- Emergent Object Storage ---
STORAGE_BASE = (os.environ.get("INTEGRATION_PROXY_URL") or "").strip() or "https://integrations.emergentagent.com"
STORAGE_URL = STORAGE_BASE.rstrip("/") + "/objstore/api/v1/storage"
EMERGENT_KEY = os.environ.get("EMERGENT_LLM_KEY")
APP_NAME = "rapportini-lavoro"
storage_key = None


def init_storage():
    global storage_key
    if storage_key:
        return storage_key
    resp = requests.post(f"{STORAGE_URL}/init", json={"emergent_key": EMERGENT_KEY}, timeout=30)
    resp.raise_for_status()
    storage_key = resp.json()["storage_key"]
    return storage_key


def put_object(path: str, data: bytes, content_type: str) -> dict:
    key = init_storage()
    resp = requests.put(
        f"{STORAGE_URL}/objects/{path}",
        headers={"X-Storage-Key": key, "Content-Type": content_type}, data=data, timeout=120,
    )
    resp.raise_for_status()
    return resp.json()


def get_object(path: str) -> tuple:
    key = init_storage()
    resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    resp.raise_for_status()
    return resp.content, resp.headers.get("Content-Type", "application/octet-stream")

app = FastAPI()
api_router = APIRouter(prefix="/api")
security = HTTPBearer(auto_error=False)

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

MONTH_NAMES_IT = ["", "Gennaio", "Febbraio", "Marzo", "Aprile", "Maggio", "Giugno",
                  "Luglio", "Agosto", "Settembre", "Ottobre", "Novembre", "Dicembre"]


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def now_utc() -> datetime:
    return datetime.now(timezone.utc)


def current_month_key() -> str:
    return now_utc().strftime("%Y-%m")


def hash_password(pw: str) -> str:
    return bcrypt.hashpw(pw.encode("utf-8")[:72], bcrypt.gensalt()).decode("utf-8")


def verify_password(pw: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(pw.encode("utf-8")[:72], hashed.encode("utf-8"))
    except Exception:
        return False


def make_token(user_id: str) -> str:
    payload = {"sub": user_id, "iat": now_utc()}
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGO)


def clean(doc: dict) -> dict:
    doc = dict(doc)
    doc.pop("_id", None)
    return doc


# ---------------------------------------------------------------------------
# Models
# ---------------------------------------------------------------------------
class RegisterIn(BaseModel):
    name: str = Field(min_length=2, max_length=80)
    email: EmailStr
    password: str = Field(min_length=6, max_length=128)


class LoginIn(BaseModel):
    email: EmailStr
    password: str


class CantiereIn(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    address: Optional[str] = ""


class ReportFields(BaseModel):
    date: str  # YYYY-MM-DD
    cantiere_id: str
    hours: float = Field(ge=0, le=24)
    drove_vehicle: bool = False
    description: str = ""
    photos: List[str] = []


class AdminReportEdit(BaseModel):
    date: str
    cantiere_id: str
    hours: float = Field(ge=0, le=24)
    drove_vehicle: bool = False
    description: str = ""
    photos: List[str] = []


class RoleIn(BaseModel):
    role: str = Field(pattern="^(employee|admin)$")


class ApproveIn(BaseModel):
    approved: bool


class ApproveAllIn(BaseModel):
    month: str


# ---------------------------------------------------------------------------
# Auth dependencies
# ---------------------------------------------------------------------------
async def get_current_user(creds: Optional[HTTPAuthorizationCredentials] = Depends(security)) -> dict:
    if creds is None:
        raise HTTPException(status_code=401, detail="Non autenticato")
    try:
        payload = jwt.decode(creds.credentials, JWT_SECRET, algorithms=[JWT_ALGO])
        user_id = payload.get("sub")
    except Exception:
        raise HTTPException(status_code=401, detail="Token non valido")
    user = await db.users.find_one({"id": user_id, "deleted_at": None})
    if not user or not user.get("approved"):
        raise HTTPException(status_code=401, detail="Account non valido o non approvato")
    return user


async def require_admin(user: dict = Depends(get_current_user)) -> dict:
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Accesso riservato agli amministratori")
    return user


def user_public(u: dict) -> dict:
    return {
        "id": u["id"],
        "name": u["name"],
        "email": u["email"],
        "role": u["role"],
        "approved": u["approved"],
    }


async def cantiere_name(cid: str) -> str:
    c = await db.cantieri.find_one({"id": cid})
    return c["name"] if c else "—"


# ---------------------------------------------------------------------------
# Report serialization
# ---------------------------------------------------------------------------
async def report_employee_view(r: dict) -> dict:
    """What the employee sees — always their own original data."""
    return {
        "id": r["id"],
        "date": r["date"],
        "month_key": r["month_key"],
        "cantiere_id": r["cantiere_id"],
        "cantiere_name": await cantiere_name(r["cantiere_id"]),
        "hours": r["hours"],
        "drove_vehicle": r["drove_vehicle"],
        "description": r["description"],
        "photos": r.get("photos", []),
        "approved": r["approved"],
    }


async def report_admin_view(r: dict) -> dict:
    """What the admin sees — admin edited version if present."""
    src = r["admin_fields"] if r.get("admin_edited") and r.get("admin_fields") else r
    photos = src.get("photos") if src.get("photos") is not None else r.get("photos", [])
    return {
        "id": r["id"],
        "user_id": r["user_id"],
        "user_name": r["user_name"],
        "date": src["date"],
        "month_key": (src["date"][:7]),
        "cantiere_id": src["cantiere_id"],
        "cantiere_name": await cantiere_name(src["cantiere_id"]),
        "hours": src["hours"],
        "drove_vehicle": src["drove_vehicle"],
        "description": src["description"],
        "photos": photos,
        "approved": r["approved"],
        "admin_edited": bool(r.get("admin_edited")),
    }


# ---------------------------------------------------------------------------
# Auth routes
# ---------------------------------------------------------------------------
@api_router.post("/auth/register")
async def register(body: RegisterIn):
    email = body.email.lower().strip()
    existing = await db.users.find_one({"email": email, "deleted_at": None})
    if existing:
        raise HTTPException(status_code=409, detail="Email già registrata")
    user = {
        "id": str(uuid.uuid4()),
        "name": body.name.strip(),
        "email": email,
        "password_hash": hash_password(body.password),
        "role": "employee",
        "approved": False,
        "created_at": now_utc().isoformat(),
        "deleted_at": None,
    }
    await db.users.insert_one(user)
    return {"message": "Registrazione completata. In attesa di approvazione dell'amministratore."}


@api_router.post("/auth/login")
async def login(body: LoginIn):
    email = body.email.lower().strip()
    user = await db.users.find_one({"email": email, "deleted_at": None})
    if not user or not verify_password(body.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Credenziali non valide")
    if not user.get("approved"):
        raise HTTPException(status_code=403, detail="Account in attesa di approvazione")
    return {"access_token": make_token(user["id"]), "user": user_public(user)}


@api_router.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return user_public(user)


# ---------------------------------------------------------------------------
# Cantieri routes
# ---------------------------------------------------------------------------
@api_router.get("/cantieri")
async def list_cantieri(user: dict = Depends(get_current_user)):
    items = await db.cantieri.find({"deleted_at": None}).sort("name", 1).to_list(1000)
    return [clean(c) for c in items]


@api_router.post("/cantieri")
async def create_cantiere(body: CantiereIn, admin: dict = Depends(require_admin)):
    c = {
        "id": str(uuid.uuid4()),
        "name": body.name.strip(),
        "address": (body.address or "").strip(),
        "created_at": now_utc().isoformat(),
        "deleted_at": None,
    }
    await db.cantieri.insert_one(c)
    return clean(c)


@api_router.patch("/cantieri/{cid}")
async def update_cantiere(cid: str, body: CantiereIn, admin: dict = Depends(require_admin)):
    res = await db.cantieri.find_one_and_update(
        {"id": cid, "deleted_at": None},
        {"$set": {"name": body.name.strip(), "address": (body.address or "").strip()}},
        return_document=True,
    )
    if not res:
        raise HTTPException(status_code=404, detail="Cantiere non trovato")
    return clean(res)


@api_router.delete("/cantieri/{cid}")
async def delete_cantiere(cid: str, admin: dict = Depends(require_admin)):
    res = await db.cantieri.find_one_and_update(
        {"id": cid, "deleted_at": None},
        {"$set": {"deleted_at": now_utc().isoformat()}},
    )
    if not res:
        raise HTTPException(status_code=404, detail="Cantiere non trovato")
    return {"message": "Cantiere eliminato"}


# ---------------------------------------------------------------------------
# Employee report routes
# ---------------------------------------------------------------------------
@api_router.get("/reports/mine")
async def my_reports(user: dict = Depends(get_current_user)):
    mk = current_month_key()
    items = await db.reports.find(
        {"user_id": user["id"], "month_key": mk, "deleted_at": None}
    ).sort("date", -1).to_list(1000)
    return [await report_employee_view(r) for r in items]


@api_router.post("/reports")
async def create_report(body: ReportFields, user: dict = Depends(get_current_user)):
    mk = body.date[:7]
    if mk != current_month_key():
        raise HTTPException(status_code=400, detail="Puoi inserire solo rapportini del mese in corso")
    c = await db.cantieri.find_one({"id": body.cantiere_id, "deleted_at": None})
    if not c:
        raise HTTPException(status_code=400, detail="Cantiere non valido")
    r = {
        "id": str(uuid.uuid4()),
        "user_id": user["id"],
        "user_name": user["name"],
        "date": body.date,
        "month_key": mk,
        "cantiere_id": body.cantiere_id,
        "hours": body.hours,
        "drove_vehicle": body.drove_vehicle,
        "description": body.description.strip(),
        "photos": body.photos,
        "approved": False,
        "admin_edited": False,
        "admin_fields": None,
        "created_at": now_utc().isoformat(),
        "updated_at": now_utc().isoformat(),
        "deleted_at": None,
    }
    await db.reports.insert_one(r)
    return await report_employee_view(r)


@api_router.patch("/reports/{rid}")
async def update_report(rid: str, body: ReportFields, user: dict = Depends(get_current_user)):
    r = await db.reports.find_one({"id": rid, "user_id": user["id"], "deleted_at": None})
    if not r:
        raise HTTPException(status_code=404, detail="Rapportino non trovato")
    if r["month_key"] != current_month_key():
        raise HTTPException(status_code=403, detail="Non puoi modificare rapportini di mesi passati")
    mk = body.date[:7]
    if mk != current_month_key():
        raise HTTPException(status_code=400, detail="La data deve essere nel mese in corso")
    c = await db.cantieri.find_one({"id": body.cantiere_id, "deleted_at": None})
    if not c:
        raise HTTPException(status_code=400, detail="Cantiere non valido")
    await db.reports.update_one({"id": rid}, {"$set": {
        "date": body.date,
        "month_key": mk,
        "cantiere_id": body.cantiere_id,
        "hours": body.hours,
        "drove_vehicle": body.drove_vehicle,
        "description": body.description.strip(),
        "photos": body.photos,
        "updated_at": now_utc().isoformat(),
    }})
    r = await db.reports.find_one({"id": rid})
    return await report_employee_view(r)


@api_router.delete("/reports/{rid}")
async def delete_report(rid: str, user: dict = Depends(get_current_user)):
    r = await db.reports.find_one({"id": rid, "user_id": user["id"], "deleted_at": None})
    if not r:
        raise HTTPException(status_code=404, detail="Rapportino non trovato")
    if r["month_key"] != current_month_key():
        raise HTTPException(status_code=403, detail="Non puoi eliminare rapportini di mesi passati")
    await db.reports.update_one({"id": rid}, {"$set": {"deleted_at": now_utc().isoformat()}})
    return {"message": "Rapportino eliminato"}


# ---------------------------------------------------------------------------
# Admin report routes
# ---------------------------------------------------------------------------
@api_router.get("/admin/reports")
async def admin_reports(
    month: Optional[str] = None,
    cantiere_id: Optional[str] = None,
    user_id: Optional[str] = None,
    admin: dict = Depends(require_admin),
):
    q: dict = {"deleted_at": None}
    if cantiere_id:
        q["cantiere_id"] = cantiere_id
    if user_id:
        q["user_id"] = user_id
    items = await db.reports.find(q).sort("date", -1).to_list(5000)
    views = [await report_admin_view(r) for r in items]
    if month:
        views = [v for v in views if v["month_key"] == month]
    return views


@api_router.patch("/admin/reports/{rid}")
async def admin_edit_report(rid: str, body: AdminReportEdit, admin: dict = Depends(require_admin)):
    r = await db.reports.find_one({"id": rid, "deleted_at": None})
    if not r:
        raise HTTPException(status_code=404, detail="Rapportino non trovato")
    c = await db.cantieri.find_one({"id": body.cantiere_id, "deleted_at": None})
    if not c:
        raise HTTPException(status_code=400, detail="Cantiere non valido")
    await db.reports.update_one({"id": rid}, {"$set": {
        "admin_edited": True,
        "admin_fields": {
            "date": body.date,
            "cantiere_id": body.cantiere_id,
            "hours": body.hours,
            "drove_vehicle": body.drove_vehicle,
            "description": body.description.strip(),
            "photos": body.photos,
        },
        "updated_at": now_utc().isoformat(),
    }})
    r = await db.reports.find_one({"id": rid})
    return await report_admin_view(r)


@api_router.patch("/admin/reports/{rid}/approve")
async def admin_approve_report(rid: str, body: ApproveIn, admin: dict = Depends(require_admin)):
    r = await db.reports.find_one_and_update(
        {"id": rid, "deleted_at": None},
        {"$set": {"approved": body.approved, "updated_at": now_utc().isoformat()}},
        return_document=True,
    )
    if not r:
        raise HTTPException(status_code=404, detail="Rapportino non trovato")
    return await report_admin_view(r)


@api_router.post("/admin/reports/approve-all")
async def approve_all(body: ApproveAllIn, admin: dict = Depends(require_admin)):
    items = await db.reports.find({"deleted_at": None, "approved": False}).to_list(10000)
    ids = []
    for r in items:
        v = await report_admin_view(r)
        if v["month_key"] == body.month:
            ids.append(r["id"])
    if ids:
        await db.reports.update_many(
            {"id": {"$in": ids}},
            {"$set": {"approved": True, "updated_at": now_utc().isoformat()}},
        )
    return {"approved": len(ids)}


@api_router.get("/admin/cantieri-summary")
async def cantieri_summary(month: Optional[str] = None, admin: dict = Depends(require_admin)):
    month = month or current_month_key()
    items = await db.reports.find({"deleted_at": None}).to_list(10000)
    cmap: Dict[str, dict] = {}
    for r in items:
        v = await report_admin_view(r)
        if v["month_key"] != month:
            continue
        cid = v["cantiere_id"]
        if cid not in cmap:
            cmap[cid] = {
                "cantiere_id": cid, "cantiere_name": v["cantiere_name"],
                "hours": 0.0, "days": set(), "reports": 0, "employees": set(),
            }
        cmap[cid]["hours"] = round(cmap[cid]["hours"] + v["hours"], 2)
        cmap[cid]["days"].add(v["date"])
        cmap[cid]["reports"] += 1
        cmap[cid]["employees"].add(v["user_id"])
    out = [
        {
            "cantiere_id": c["cantiere_id"], "cantiere_name": c["cantiere_name"],
            "hours": round(c["hours"], 2), "days": len(c["days"]),
            "reports": c["reports"], "employees": len(c["employees"]),
        }
        for c in cmap.values()
    ]
    out.sort(key=lambda x: -x["hours"])
    return {"month": month, "cantieri": out}


async def build_matrix(month: str) -> dict:
    year, mon = int(month[:4]), int(month[5:7])
    first_next = date(year + (1 if mon == 12 else 0), (mon % 12) + 1, 1)
    days_in_month = (first_next - date(year, mon, 1)).days
    items = await db.reports.find({"deleted_at": None}).to_list(10000)
    emap: Dict[str, dict] = {}
    for r in items:
        v = await report_admin_view(r)
        if v["month_key"] != month:
            continue
        uid = v["user_id"]
        if uid not in emap:
            emap[uid] = {"user_id": uid, "name": v["user_name"], "daily": {}, "total": 0.0}
        try:
            day = int(v["date"][8:10])
        except Exception:
            continue
        emap[uid]["daily"][str(day)] = round(emap[uid]["daily"].get(str(day), 0.0) + v["hours"], 2)
        emap[uid]["total"] = round(emap[uid]["total"] + v["hours"], 2)
    employees = sorted(emap.values(), key=lambda e: e["name"].lower())
    return {
        "month": month,
        "days": list(range(1, days_in_month + 1)),
        "employees": employees,
    }


@api_router.get("/admin/matrix")
async def matrix(month: Optional[str] = None, admin: dict = Depends(require_admin)):
    month = month or current_month_key()
    return await build_matrix(month)


def month_label(month: str) -> str:
    y, m = int(month[:4]), int(month[5:7])
    return f"{MONTH_NAMES_IT[m]} {y}"


@api_router.get("/admin/export")
async def export_matrix(month: Optional[str] = None, format: str = "excel", admin: dict = Depends(require_admin)):
    month = month or current_month_key()
    data = await build_matrix(month)
    days = data["days"]
    employees = data["employees"]
    label = month_label(month)

    if format == "excel":
        wb = Workbook()
        ws = wb.active
        ws.title = "Rapportino"
        header = ["Dipendente"] + [str(d) for d in days] + ["Totale"]
        ws.append(header)
        head_fill = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center")
        for e in employees:
            row = [e["name"]] + [e["daily"].get(str(d), "") for d in days] + [e["total"]]
            ws.append(row)
        ws.column_dimensions["A"].width = 24
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        b64 = base64.b64encode(buf.read()).decode("utf-8")
        return {
            "filename": f"rapportino_{month}.xlsx",
            "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "base64": b64,
        }

    # PDF
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * mm,
                            rightMargin=10 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
    styles = getSampleStyleSheet()
    elements = [Paragraph(f"Rapportino Ore — {label}", styles["Title"]), Spacer(1, 8)]
    header = ["Dipendente"] + [str(d) for d in days] + ["Tot"]
    table_data = [header]
    for e in employees:
        row = [e["name"]] + [str(e["daily"].get(str(d), "")) for d in days] + [str(e["total"])]
        table_data.append(row)
    if len(table_data) == 1:
        table_data.append(["Nessun dato"] + [""] * (len(days) + 1))
    t = Table(table_data, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0284C7")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 6),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D1D5DB")),
        ("ALIGN", (1, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
    ]))
    elements.append(t)
    doc.build(elements)
    buf.seek(0)
    b64 = base64.b64encode(buf.read()).decode("utf-8")
    return {
        "filename": f"rapportino_{month}.pdf",
        "mime": "application/pdf",
        "base64": b64,
    }


# ---------------------------------------------------------------------------
# Admin user management
# ---------------------------------------------------------------------------
@api_router.get("/admin/users")
async def admin_users(admin: dict = Depends(require_admin)):
    users = await db.users.find({"deleted_at": None}).sort("created_at", -1).to_list(2000)
    return [user_public(u) for u in users]


@api_router.patch("/admin/users/{uid}/approve")
async def approve_user(uid: str, body: ApproveIn, admin: dict = Depends(require_admin)):
    res = await db.users.find_one_and_update(
        {"id": uid, "deleted_at": None},
        {"$set": {"approved": body.approved}},
        return_document=True,
    )
    if not res:
        raise HTTPException(status_code=404, detail="Utente non trovato")
    return user_public(res)


@api_router.patch("/admin/users/{uid}/role")
async def set_role(uid: str, body: RoleIn, admin: dict = Depends(require_admin)):
    if uid == admin["id"] and body.role != "admin":
        raise HTTPException(status_code=400, detail="Non puoi rimuovere il tuo ruolo admin")
    res = await db.users.find_one_and_update(
        {"id": uid, "deleted_at": None},
        {"$set": {"role": body.role}},
        return_document=True,
    )
    if not res:
        raise HTTPException(status_code=404, detail="Utente non trovato")
    return user_public(res)


@api_router.delete("/admin/users/{uid}")
async def delete_user(uid: str, admin: dict = Depends(require_admin)):
    if uid == admin["id"]:
        raise HTTPException(status_code=400, detail="Non puoi eliminare te stesso")
    res = await db.users.find_one_and_update(
        {"id": uid, "deleted_at": None},
        {"$set": {"deleted_at": now_utc().isoformat(), "approved": False}},
    )
    if not res:
        raise HTTPException(status_code=404, detail="Utente non trovato")
    return {"message": "Utente eliminato"}


# ---------------------------------------------------------------------------
# File upload / download (Emergent Object Storage)
# ---------------------------------------------------------------------------
@api_router.post("/upload")
async def upload_file(file: UploadFile = File(...), user: dict = Depends(get_current_user)):
    ext = "jpg"
    if file.filename and "." in file.filename:
        ext = file.filename.rsplit(".", 1)[-1].lower()[:5] or "jpg"
    key = f"{APP_NAME}/uploads/{user['id']}/{uuid.uuid4()}.{ext}"
    data = await file.read()
    try:
        await run_in_threadpool(put_object, key, data, file.content_type or "image/jpeg")
    except requests.HTTPError as e:
        status_code = e.response.status_code if e.response is not None else 500
        if status_code == 402:
            raise HTTPException(status_code=402, detail="Spazio di archiviazione esaurito")
        raise HTTPException(status_code=502, detail="Errore caricamento file")
    return {"path": key}


@api_router.get("/files/{path:path}")
async def get_file(
    path: str,
    token: Optional[str] = Query(default=None),
    creds: Optional[HTTPAuthorizationCredentials] = Depends(security),
):
    raw = token or (creds.credentials if creds else None)
    if not raw:
        raise HTTPException(status_code=401, detail="Non autenticato")
    try:
        payload = jwt.decode(raw, JWT_SECRET, algorithms=[JWT_ALGO])
        user_id = payload.get("sub")
    except Exception:
        raise HTTPException(status_code=401, detail="Token non valido")
    user = await db.users.find_one({"id": user_id, "deleted_at": None})
    if not user or not user.get("approved"):
        raise HTTPException(status_code=401, detail="Account non valido")
    parts = path.split("/")
    owner = parts[2] if len(parts) >= 3 else None
    if user.get("role") != "admin" and owner != user_id:
        raise HTTPException(status_code=403, detail="Accesso negato")
    try:
        content, ctype = await run_in_threadpool(get_object, path)
    except Exception:
        raise HTTPException(status_code=404, detail="File non trovato")
    return Response(content=content, media_type=ctype)


@api_router.get("/")
async def root():
    return {"message": "Rapportini API"}


app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
async def seed_admin():
    try:
        await run_in_threadpool(init_storage)
        logger.info("Object storage initialized")
    except Exception as e:
        logger.warning(f"Storage init failed: {e}")
    await db.users.create_index("email")
    await db.users.create_index("id")
    await db.reports.create_index("user_id")
    existing = await db.users.find_one({"role": "admin", "deleted_at": None})
    if not existing:
        admin = {
            "id": str(uuid.uuid4()),
            "name": "Amministratore",
            "email": "admin@rapportini.it",
            "password_hash": hash_password("Admin1234!"),
            "role": "admin",
            "approved": True,
            "created_at": now_utc().isoformat(),
            "deleted_at": None,
        }
        await db.users.insert_one(admin)
        logger.info("Seeded default admin admin@rapportini.it")


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
